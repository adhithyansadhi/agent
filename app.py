from gevent import monkey
monkey.patch_all()

import os
import requests
import pandas as pd
import time
import threading
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
import uuid
import queue
import tempfile
import shutil
from pathlib import Path
import traceback # <--- IMPORT THIS MODULE
from flask import Flask, render_template, request, jsonify, Response, send_file, after_this_request

# --- Basic Flask App Setup ---
app = Flask(__name__)
load_dotenv()

# In-memory store for job progress and logs
JOBS = {}

# --- Configuration ---
API_KEY = os.getenv("GEMINI_API_KEY")
API_URL = "https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent"
DELAY_BETWEEN_FILES = 2

# ==============================================================================
# 1. CORE LOGIC (No changes in this section)
# ==============================================================================
# ... (All functions like get_key_column_from_ai, compare_dataframes, etc., remain unchanged)
def get_key_column_from_ai(df1_headers, df2_headers, log_queue):
    system_prompt = """
    You are a data analysis expert. Your task is to identify the single best column for joining two datasets based on their headers.
    Follow these rules strictly:
    1.  **Prioritize descriptive name columns** (like 'ProductName', 'Policy Name') over generic ID columns (like 'ProductID', 'ID').
    2.  The chosen column must exist in both lists of headers provided.
    3.  Respond with ONLY the column name as a single, unformatted string (e.g., "ProductName").
    """
    user_query = f"File 1 Headers: {df1_headers}\nFile 2 Headers: {df2_headers}\n\nBased on the rules, what is the single best column name to join these files on?"
    headers = {"Content-Type": "application/json"}
    payload = {
        "contents": [{"parts": [{"text": user_query}]}],
        "systemInstruction": {"parts": [{"text": system_prompt}]},
        "generationConfig": {"responseMimeType": "text/plain"}
    }
    try:
        response = requests.post(f"{API_URL}?key={API_KEY}", headers=headers, json=payload)
        response.raise_for_status()
        return response.json()['candidates'][0]['content']['parts'][0]['text'].strip()
    except requests.exceptions.RequestException as e:
        log_queue.put(f"API request failed to get key column: {e}")
        return None

def compare_dataframes(df1, df2, key_column):
    merged_df = pd.merge(df1, df2, on=key_column, how='outer', suffixes=('_1', '_2'))
    results = []
    value_cols_1 = [col for col in df1.columns if col != key_column and pd.api.types.is_numeric_dtype(df1[col])]
    for _, row in merged_df.iterrows():
        record_id = row[key_column]
        metrics = []
        for col in value_cols_1:
            val1_col, val2_col = f"{col}_1", f"{col}_2"
            value1, value2 = row[val1_col], row[val2_col]
            status, percentage_change = "", None
            if pd.isna(value1):
                status = "New"; value1 = None
            elif pd.isna(value2):
                status = "Missing"; value2 = None
            else:
                if value1 == value2:
                    status = "Unchanged"; percentage_change = 0.0
                else:
                    status = "Updated"
                    if value1 == 0: percentage_change = 100.0 if value2 != 0 else 0.0
                    else: percentage_change = ((value2 - value1) / value1) * 100
            metrics.append({"metric_name": col, "value1": value1, "value2": value2, "percentage_change": percentage_change, "status": status})
        results.append({"id": record_id, "metrics": metrics})
    return results

def save_styled_results_to_excel(data, filename, output_folder, log_queue):
    if not data:
        log_queue.put(f"No data to save for {filename}.")
        return
    wb = Workbook()
    ws = wb.active
    grey_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    pink_fill = PatternFill(start_color="FFC0CB", end_color="FFC0CB", fill_type="solid")
    green_font = Font(color="008000", bold=True)
    red_font = Font(color="FF0000", bold=True)
    header_font = Font(bold=True)
    center_align = Alignment(horizontal='center', vertical='center')
    first_record = next((r for r in data if r.get('metrics')), None)
    if not first_record:
        log_queue.put(f"No metrics found to create a file for {filename}.")
        return
    metric_names = [m['metric_name'] for m in first_record['metrics']]
    header = ["ID"] + metric_names
    ws.append(header)
    for cell in ws[1]:
        cell.font = header_font; cell.alignment = center_align
    for record in data:
        row_to_append = [record.get('id')]
        metrics_by_name = {m['metric_name']: m for m in record.get('metrics', [])}
        for name in metric_names:
            metric = metrics_by_name.get(name, {}); status = metric.get('status'); cell_value_str = ""
            if status == "Updated":
                v1, v2, chg = metric.get('value1'), metric.get('value2'), metric.get('percentage_change')
                sign = "+" if isinstance(chg, (int, float)) and chg > 0 else ""
                cell_value_str = f"{v1} → {v2} ({sign}{chg:.2f}%)" if isinstance(chg, (int, float)) else f"{v1} → {v2}"
            elif status in ("New", "Missing", "Unchanged"):
                val = metric.get('value2') if status == "New" else metric.get('value1')
                cell_value_str = str(val) if val is not None else "N/A"
            row_to_append.append(cell_value_str)
        ws.append(row_to_append)
        current_row = ws.max_row
        for col_idx, name in enumerate(metric_names, 2):
            cell = ws.cell(row=current_row, column=col_idx)
            metric = metrics_by_name.get(name, {}); status = metric.get('status')
            if status == "Updated":
                chg = metric.get('percentage_change')
                if isinstance(chg, (int, float)): cell.font = green_font if chg > 0 else red_font
            elif status == "New": cell.fill = pink_fill
            elif status == "Missing": cell.fill = grey_fill
            cell.alignment = center_align
    for column_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in column_cells if cell.value)
        ws.column_dimensions[column_cells[0].column_letter].width = length + 5

    output_filename = f"comparison_{os.path.basename(filename)}"
    output_filepath = os.path.join(output_folder, os.path.splitext(output_filename)[0] + ".xlsx")
    try:
        wb.save(output_filepath)
        log_queue.put(f"Successfully saved styled results to {output_filename}")
    except Exception as e:
        log_queue.put(f"Failed to save Excel file: {e}")

# ==============================================================================
# 2. BACKGROUND TASK (MODIFIED FOR DEBUGGING)
# ==============================================================================

def run_comparison_background_task(job_id, files1_map, files2_map, output_name, temp_dir1, temp_dir2):
    log_queue = JOBS[job_id]['queue']
    output_folder = None
    try:
        # ... (The main logic of the function remains the same)
        if not API_KEY:
            log_queue.put("CRITICAL ERROR: GEMINI_API_KEY not found in .env file.")
            JOBS[job_id]['status'] = 'failed'
            return

        output_folder = tempfile.mkdtemp()
        log_queue.put(f"Created temporary output directory.")

        common_keys = sorted(list(set(files1_map.keys()).intersection(files2_map.keys())))
        if not common_keys:
            log_queue.put("No common Excel files (.xlsx) found in the selected folders.")
            JOBS[job_id]['status'] = 'completed'
            log_queue.put(f"DONE:{job_id}")
            return

        log_queue.put(f"Found {len(common_keys)} common files to compare.")
        for i, key in enumerate(common_keys):
            log_queue.put(f"\n--- [{i+1}/{len(common_keys)}] Processing: {key} ---")
            filepath1 = files1_map[key]
            filepath2 = files2_map[key]
            df1, df2 = pd.read_excel(filepath1), pd.read_excel(filepath2)
            log_queue.put("Contacting AI to identify the key column...")
            key_column = get_key_column_from_ai(list(df1.columns), list(df2.columns), log_queue)
            if not key_column or key_column not in df1.columns or key_column not in df2.columns:
                log_queue.put(f"Error: Could not determine a valid key column. AI returned: '{key_column}'. Skipping file.")
                continue
            log_queue.put(f"AI identified key column: '{key_column}'")
            log_queue.put("Performing local comparison...")
            comparison_data = compare_dataframes(df1, df2, key_column)
            log_queue.put("Saving styled Excel report...")
            save_styled_results_to_excel(comparison_data, os.path.basename(key), output_folder, log_queue)
            if i < len(common_keys) - 1:
                log_queue.put(f"Waiting for {DELAY_BETWEEN_FILES} seconds...")
                time.sleep(DELAY_BETWEEN_FILES)

        log_queue.put("\n--- All comparisons complete! Zipping results... ---")
        zip_path_base = os.path.join(tempfile.gettempdir(), output_name)
        zip_path = shutil.make_archive(zip_path_base, 'zip', output_folder)
        JOBS[job_id]['zip_path'] = zip_path
        JOBS[job_id]['status'] = 'completed'
        log_queue.put(f"DONE:{job_id}")

    except Exception as e:
        # ### THIS IS THE DEBUGGING FIX ###
        # In production, you might want to log this to a file instead.
        # For development, we send the full traceback to the frontend.
        error_details = traceback.format_exc()
        log_queue.put("\n" + "="*20 + " SERVER ERROR " + "="*20)
        log_queue.put(f"A critical error occurred: {e}")
        log_queue.put("Full Traceback:")
        log_queue.put(error_details)
        log_queue.put("="*54)
        JOBS[job_id]['status'] = 'failed'
        # ### END OF FIX ###

    finally:
        final_status = JOBS[job_id].get('status', 'unknown')
        log_queue.put(f"DONE:{final_status}:{job_id}")
        # This block will run whether there was an error or not.
        # We must signal the frontend that the process is over.
        # Clean up the temporary directories
        shutil.rmtree(temp_dir1)
        shutil.rmtree(temp_dir2)
        if output_folder:
            shutil.rmtree(output_folder)

# ==============================================================================
# 3. FLASK ROUTES (No changes in this section)
# ==============================================================================
# ... (All routes like @app.route('/'), @app.route('/compare'), etc., remain unchanged)
@app.route('/')
def index():
    return render_template('index.html')

@app.route('/compare', methods=['POST'])
def start_comparison():
    job_id = str(uuid.uuid4())
    JOBS[job_id] = {'queue': queue.Queue(), 'status': 'running'}

    try:
        files1 = request.files.getlist('folder1_files')
        files2 = request.files.getlist('folder2_files')
        output_name = request.form.get('output_name', 'comparison_results')

        if not files1 or not files2:
            return jsonify({'error': 'Please select files for both folders.'}), 400

        temp_dir1 = tempfile.mkdtemp()
        temp_dir2 = tempfile.mkdtemp()

        files1_map = {}
        files2_map = {}

        for file in files1:
            if file.filename:
                p = Path(file.filename)
                key = os.path.join(*p.parts[1:]) if len(p.parts) > 1 else p.name
                dest_path = os.path.join(temp_dir1, file.filename)
                os.makedirs(os.path.dirname(dest_path), exist_ok=True)
                file.save(dest_path)
                files1_map[key] = dest_path

        for file in files2:
            if file.filename:
                p = Path(file.filename)
                key = os.path.join(*p.parts[1:]) if len(p.parts) > 1 else p.name
                dest_path = os.path.join(temp_dir2, file.filename)
                os.makedirs(os.path.dirname(dest_path), exist_ok=True)
                file.save(dest_path)
                files2_map[key] = dest_path
        
        thread = threading.Thread(
            target=run_comparison_background_task,
            args=(job_id, files1_map, files2_map, output_name, temp_dir1, temp_dir2)
        )
        thread.daemon = True
        thread.start()

        return jsonify({'job_id': job_id})

    except Exception as e:
        JOBS[job_id]['status'] = 'failed'
        if 'temp_dir1' in locals() and os.path.exists(temp_dir1):
             shutil.rmtree(temp_dir1)
        if 'temp_dir2' in locals() and os.path.exists(temp_dir2):
             shutil.rmtree(temp_dir2)
        return jsonify({'error': str(e)}), 500

@app.route('/stream/<job_id>')
def stream_logs(job_id):
    def generate():
        if job_id not in JOBS:
            yield f"data: Error: Job ID {job_id} not found.\n\n"
            return
        
        log_queue = JOBS[job_id]['queue']
        while True:
            try:
                message = log_queue.get(timeout=5)
                if message.startswith(f"DONE:{job_id}"):
                    # Send the final DONE message before breaking
                    yield f"data: {message}\n\n"
                    break
                yield f"data: {message}\n\n"
            except queue.Empty:
                if JOBS[job_id]['status'] != 'running':
                    break
                yield ": heartbeat\n\n"
    
    return Response(generate(), mimetype='text/event-stream')

@app.route('/download/<job_id>')
def download_results(job_id):
    if job_id in JOBS and JOBS[job_id]['status'] == 'completed':
        zip_path = JOBS[job_id].get('zip_path')
        if zip_path and os.path.exists(zip_path):
            
            @after_this_request
            def cleanup(response):
                try:
                    os.remove(zip_path)
                    del JOBS[job_id]
                except Exception as e:
                    print(f"Error during cleanup for job {job_id}: {e}")
                return response
            
            download_name = f"{os.path.basename(zip_path)}"
            return send_file(zip_path, as_attachment=True, download_name=download_name)
    
    return "Error: File not found or job not complete.", 404

# if __name__ == "__main__":
#     app.run(debug=True)